import io
import fitz  # PyMuPDF
import docx
from fastapi import UploadFile, HTTPException

MAX_FILE_SIZE_MB = 5
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx"}


async def extract_text_from_resume(file: UploadFile) -> dict:
    """
    Validates file type/size and extracts plain text from PDF or DOCX files.
    """
    filename = file.filename or ""
    file_ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    # 1. Validate file extension
    if file_ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file format '{file_ext}'. Please upload a .PDF or .DOCX file."
        )

    # 2. Read file bytes and validate file size
    file_bytes = await file.read()
    file_size = len(file_bytes)

    if file_size == 0:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    if file_size > MAX_FILE_SIZE_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"File size exceeds the {MAX_FILE_SIZE_MB}MB limit."
        )

    # 3. Extract text based on file extension
    extracted_text = ""
    try:
        if file_ext == ".pdf":
            extracted_text = _parse_pdf(file_bytes)
        elif file_ext == ".docx":
            extracted_text = _parse_docx(file_bytes)
    except Exception as e:
        raise HTTPException(
            status_code=422,
            detail=f"Failed to read the {file_ext.upper()} file. Ensure it is not corrupted or password-protected. ({str(e)})"
        )

    cleaned_text = extracted_text.strip()
    if len(cleaned_text) < 30:
        raise HTTPException(
            status_code=422,
            detail="Could not extract enough readable text from the resume. If this is a scanned image PDF, please upload a text-based PDF or DOCX."
        )

    return {
        "filename": filename,
        "file_size_bytes": file_size,
        "file_size_kb": round(file_size / 1024, 2),
        "extracted_text": cleaned_text,
    }


def _parse_pdf(file_bytes: bytes) -> str:
    """Extracts text from PDF bytes using PyMuPDF."""
    text_pages = []
    with fitz.open(stream=file_bytes, filetype="pdf") as pdf_doc:
        for page in pdf_doc:
            text_pages.append(page.get_text("text"))
    return "\n".join(text_pages)


def _parse_docx(file_bytes: bytes) -> str:
    """Extracts text from DOCX bytes using python-docx (paragraphs + tables)."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also extract text from tables in case the resume uses a table layout
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs.append(" | ".join(row_text))

    return "\n".join(paragraphs)